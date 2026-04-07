from fastapi import FastAPI, APIRouter, HTTPException, UploadFile, File
from fastapi.responses import StreamingResponse
from dotenv import load_dotenv
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
import os
import logging
from pathlib import Path
from pydantic import BaseModel, Field
from typing import List, Optional
import uuid
from datetime import datetime, timedelta
import re
import hashlib
import io
import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

# MongoDB connection
mongo_url = os.environ['MONGO_URL']
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ['DB_NAME']]

# Create the main app without a prefix
app = FastAPI()

# Create a router with the /api prefix
api_router = APIRouter(prefix="/api")

# Default categories (will be used to seed if no categories exist)
DEFAULT_CATEGORIES = [
    {"id": "food", "name": "Food", "icon": "fast-food", "color": "#FF6B6B", "is_default": True},
    {"id": "transport", "name": "Transport", "icon": "car", "color": "#4ECDC4", "is_default": True},
    {"id": "shopping", "name": "Shopping", "icon": "cart", "color": "#45B7D1", "is_default": True},
    {"id": "bills", "name": "Bills", "icon": "receipt", "color": "#96CEB4", "is_default": True},
    {"id": "entertainment", "name": "Entertainment", "icon": "game-controller", "color": "#DDA0DD", "is_default": True},
    {"id": "healthcare", "name": "Healthcare", "icon": "medical", "color": "#98D8C8", "is_default": True},
    {"id": "income", "name": "Income", "icon": "cash", "color": "#2ECC71", "is_default": True},
    {"id": "others", "name": "Others", "icon": "ellipsis-horizontal", "color": "#95A5A6", "is_default": True}
]

# Define Models
class Expense(BaseModel):
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    amount: float
    category_id: str
    description: str
    merchant: Optional[str] = None
    transaction_type: str = "debit"  # debit or credit
    source: str = "manual"  # manual, sms, or pdf
    sms_body: Optional[str] = None
    transaction_date: Optional[datetime] = None
    fingerprint: Optional[str] = None  # For duplicate detection
    created_at: datetime = Field(default_factory=datetime.utcnow)

class ExpenseCreate(BaseModel):
    amount: float
    category_id: str
    description: str
    merchant: Optional[str] = None
    transaction_type: str = "debit"
    source: str = "manual"
    sms_body: Optional[str] = None
    transaction_date: Optional[str] = None

class Category(BaseModel):
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    name: str
    icon: str
    color: str
    is_default: bool = False
    created_at: datetime = Field(default_factory=datetime.utcnow)

class CategoryCreate(BaseModel):
    name: str
    icon: str
    color: str

class KeywordRule(BaseModel):
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    keyword: str
    category_id: str
    created_at: datetime = Field(default_factory=datetime.utcnow)

class KeywordRuleCreate(BaseModel):
    keyword: str
    category_id: str

class SMSParseRequest(BaseModel):
    sms_body: str

class SMSParseResponse(BaseModel):
    success: bool
    amount: Optional[float] = None
    transaction_type: Optional[str] = None
    merchant: Optional[str] = None
    suggested_category_id: Optional[str] = None
    description: Optional[str] = None
    transaction_date: Optional[str] = None
    is_duplicate: bool = False
    error: Optional[str] = None

# Helper function to generate fingerprint for duplicate detection
def generate_fingerprint(amount: float, transaction_type: str, transaction_date: Optional[datetime] = None) -> str:
    """Generate a fingerprint based on amount, type, and date for duplicate detection"""
    date_str = transaction_date.strftime("%Y-%m-%d") if transaction_date else datetime.utcnow().strftime("%Y-%m-%d")
    fingerprint_str = f"{amount:.2f}|{transaction_type}|{date_str}"
    return hashlib.md5(fingerprint_str.encode()).hexdigest()

# Helper function to check for duplicates
async def check_duplicate(amount: float, transaction_type: str, transaction_date: Optional[datetime] = None) -> bool:
    """Check if a similar transaction already exists"""
    fingerprint = generate_fingerprint(amount, transaction_type, transaction_date)
    existing = await db.expenses.find_one({"fingerprint": fingerprint})
    return existing is not None

# Helper function to parse SMS
async def parse_sms_message(sms_body: str, check_dup: bool = True) -> dict:
    """Parse SMS message to extract transaction details"""
    result = {
        "success": False,
        "amount": None,
        "transaction_type": None,
        "merchant": None,
        "suggested_category_id": "others",
        "description": None,
        "transaction_date": None,
        "is_duplicate": False,
        "error": None
    }
    
    sms_lower = sms_body.lower()
    
    # Detect transaction type
    if "debited" in sms_lower or "debit" in sms_lower or "spent" in sms_lower or "paid" in sms_lower or "withdrawn" in sms_lower or "send" in sms_lower or "sent" in sms_lower or "transferred" in sms_lower:
        result["transaction_type"] = "debit"
    elif "credited" in sms_lower or "credit" in sms_lower or "received" in sms_lower or "deposited" in sms_lower:
        result["transaction_type"] = "credit"
    else:
        result["error"] = "Could not detect transaction type (Debited/Credited/Send keywords not found)"
        return result
    
    # Extract amount - various patterns
    amount_patterns = [
        r'(?:rs\.?|inr|₹)\s*([\d,]+(?:\.\d{2})?)',
        r'([\d,]+(?:\.\d{2})?)\s*(?:rs\.?|inr|₹)',
        r'(?:amount|amt)\s*(?:of)?\s*(?:rs\.?|inr|₹)?\s*([\d,]+(?:\.\d{2})?)',
        r'(?:debited|credited|spent|paid|received|send|sent|transferred)\s*(?:for)?\s*(?:rs\.?|inr|₹)?\s*([\d,]+(?:\.\d{2})?)',
        r'([\d,]+(?:\.\d{2})?)\s*(?:has been|is|was)\s*(?:debited|credited)',
    ]
    
    amount = None
    for pattern in amount_patterns:
        match = re.search(pattern, sms_lower)
        if match:
            amount_str = match.group(1).replace(',', '')
            try:
                amount = float(amount_str)
                break
            except ValueError:
                continue
    
    if amount is None:
        result["error"] = "Could not extract amount from SMS"
        return result
    
    result["amount"] = amount
    
    # Extract date
    date_patterns = [
        r'(\d{1,2}[/-]\d{1,2}[/-]\d{2,4})',
        r'(\d{1,2}\s+(?:jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*\s+\d{2,4})',
    ]
    
    for pattern in date_patterns:
        match = re.search(pattern, sms_lower)
        if match:
            result["transaction_date"] = match.group(1)
            break
    
    # Extract merchant/recipient
    merchant_patterns = [
        r'(?:to|at|for|from)\s+([A-Za-z0-9\s]+?)(?:\s+on|\s+ref|\s+avl|\.)',
        r'(?:merchant|vendor|shop|store)[:\s]+([A-Za-z0-9\s]+)',
        r'(?:UPI|IMPS|NEFT)[:\s]*([A-Za-z0-9@\s]+?)(?:\s|$)',
    ]
    
    for pattern in merchant_patterns:
        match = re.search(pattern, sms_body, re.IGNORECASE)
        if match:
            merchant = match.group(1).strip()
            merchant = re.sub(r'\s+', ' ', merchant)
            if len(merchant) > 2:
                result["merchant"] = merchant[:50]
                break
    
    # Get keyword rules for category suggestion
    keyword_rules = await db.keyword_rules.find().to_list(100)
    
    # Check if any keyword matches
    for rule in keyword_rules:
        if rule['keyword'].lower() in sms_lower:
            result["suggested_category_id"] = rule['category_id']
            break
    
    # If no custom rule matched, try default keyword matching
    if result["suggested_category_id"] == "others":
        default_keywords = {
            "food": ["swiggy", "zomato", "restaurant", "food", "cafe", "dominos", "mcdonalds", "kfc", "pizza", "burger"],
            "transport": ["uber", "ola", "rapido", "metro", "fuel", "petrol", "diesel", "parking", "toll"],
            "shopping": ["amazon", "flipkart", "myntra", "ajio", "nykaa", "meesho", "mall", "store"],
            "bills": ["electricity", "water", "gas", "internet", "broadband", "mobile", "recharge", "bill", "jio", "airtel"],
            "entertainment": ["netflix", "spotify", "hotstar", "prime", "movie", "ticket", "bookmyshow", "pvr", "inox"],
            "healthcare": ["pharmacy", "hospital", "clinic", "medical", "doctor", "medicine", "apollo", "1mg", "pharmeasy"],
            "income": ["salary", "refund", "cashback", "reward", "bonus"]
        }
        
        for category_id, keywords in default_keywords.items():
            for keyword in keywords:
                if keyword in sms_lower:
                    result["suggested_category_id"] = category_id
                    break
            if result["suggested_category_id"] != "others":
                break
    
    # Generate description
    if result["transaction_type"] == "credit":
        result["description"] = f"Received from {result['merchant']}" if result["merchant"] else "Amount credited"
    else:
        result["description"] = f"Payment to {result['merchant']}" if result["merchant"] else "Amount debited"
    
    # Check for duplicates
    if check_dup:
        trans_date = None
        if result["transaction_date"]:
            try:
                # Try to parse the date
                for fmt in ["%d/%m/%y", "%d-%m-%y", "%d/%m/%Y", "%d-%m-%Y"]:
                    try:
                        trans_date = datetime.strptime(result["transaction_date"], fmt)
                        break
                    except:
                        continue
            except:
                pass
        
        is_dup = await check_duplicate(result["amount"], result["transaction_type"], trans_date)
        result["is_duplicate"] = is_dup
    
    result["success"] = True
    return result

# Parse PDF bank statement
async def parse_pdf_statement(pdf_content: bytes) -> List[dict]:
    """Parse PDF bank statement and extract transactions"""
    transactions = []
    
    try:
        with pdfplumber.open(io.BytesIO(pdf_content)) as pdf:
            for page in pdf.pages:
                text = page.extract_text() or ""
                tables = page.extract_tables()
                
                # Try to extract from tables first
                for table in tables:
                    for row in table:
                        if row and len(row) >= 3:
                            # Look for rows with amounts
                            row_text = ' '.join([str(cell) if cell else '' for cell in row])
                            
                            # Try to find amount patterns
                            amount_match = re.search(r'([\d,]+(?:\.\d{2})?)', row_text)
                            if amount_match:
                                try:
                                    amount = float(amount_match.group(1).replace(',', ''))
                                    if amount > 0:
                                        # Determine if debit or credit
                                        row_lower = row_text.lower()
                                        if any(kw in row_lower for kw in ['dr', 'debit', 'withdrawal', 'paid', 'transfer to']):
                                            trans_type = "debit"
                                        elif any(kw in row_lower for kw in ['cr', 'credit', 'deposit', 'received', 'transfer from']):
                                            trans_type = "credit"
                                        else:
                                            continue
                                        
                                        # Extract date
                                        date_match = re.search(r'(\d{1,2}[/-]\d{1,2}[/-]\d{2,4})', row_text)
                                        trans_date = date_match.group(1) if date_match else None
                                        
                                        transactions.append({
                                            "amount": amount,
                                            "transaction_type": trans_type,
                                            "description": row_text[:100],
                                            "transaction_date": trans_date
                                        })
                                except:
                                    continue
                
                # Also try to parse from text patterns
                lines = text.split('\n')
                for line in lines:
                    # Look for transaction patterns in text
                    if re.search(r'(?:debit|credit|dr|cr|withdrawal|deposit)', line.lower()):
                        amount_match = re.search(r'([\d,]+(?:\.\d{2})?)', line)
                        if amount_match:
                            try:
                                amount = float(amount_match.group(1).replace(',', ''))
                                if amount > 0 and amount < 10000000:  # Reasonable limit
                                    line_lower = line.lower()
                                    if any(kw in line_lower for kw in ['dr', 'debit', 'withdrawal']):
                                        trans_type = "debit"
                                    elif any(kw in line_lower for kw in ['cr', 'credit', 'deposit']):
                                        trans_type = "credit"
                                    else:
                                        continue
                                    
                                    date_match = re.search(r'(\d{1,2}[/-]\d{1,2}[/-]\d{2,4})', line)
                                    trans_date = date_match.group(1) if date_match else None
                                    
                                    # Avoid duplicates in same PDF
                                    fingerprint = f"{amount}|{trans_type}|{trans_date}"
                                    if not any(t.get('_fp') == fingerprint for t in transactions):
                                        transactions.append({
                                            "amount": amount,
                                            "transaction_type": trans_type,
                                            "description": line[:100].strip(),
                                            "transaction_date": trans_date,
                                            "_fp": fingerprint
                                        })
                            except:
                                continue
    except Exception as e:
        logging.error(f"PDF parsing error: {e}")
    
    # Remove internal fingerprint
    for t in transactions:
        t.pop('_fp', None)
    
    return transactions

# Initialize categories on startup
@app.on_event("startup")
async def init_categories():
    """Initialize default categories if none exist"""
    count = await db.categories.count_documents({})
    if count == 0:
        for cat in DEFAULT_CATEGORIES:
            await db.categories.insert_one(cat)
        logging.info("Initialized default categories")

# Routes
@api_router.get("/")
async def root():
    return {"message": "Expense Tracker API"}

# Categories CRUD
@api_router.get("/categories", response_model=List[Category])
async def get_categories():
    categories = await db.categories.find().to_list(100)
    if not categories:
        # Return defaults if empty
        return [Category(**cat) for cat in DEFAULT_CATEGORIES]
    return [Category(**cat) for cat in categories]

@api_router.post("/categories", response_model=Category)
async def create_category(category: CategoryCreate):
    # Check if name already exists
    existing = await db.categories.find_one({"name": {"$regex": f"^{category.name}$", "$options": "i"}})
    if existing:
        raise HTTPException(status_code=400, detail="Category with this name already exists")
    
    cat_dict = category.dict()
    cat_obj = Category(**cat_dict)
    await db.categories.insert_one(cat_obj.dict())
    return cat_obj

@api_router.delete("/categories/{category_id}")
async def delete_category(category_id: str):
    # Check if it's a default category
    cat = await db.categories.find_one({"id": category_id})
    if cat and cat.get("is_default"):
        raise HTTPException(status_code=400, detail="Cannot delete default category")
    
    # Check if any expenses use this category
    expense_count = await db.expenses.count_documents({"category_id": category_id})
    if expense_count > 0:
        raise HTTPException(status_code=400, detail=f"Cannot delete category with {expense_count} expenses. Reassign expenses first.")
    
    result = await db.categories.delete_one({"id": category_id})
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Category not found")
    return {"message": "Category deleted successfully"}

# Expenses CRUD
@api_router.post("/expenses", response_model=Expense)
async def create_expense(expense: ExpenseCreate):
    expense_dict = expense.dict()
    
    # Parse transaction date if provided
    trans_date = None
    if expense_dict.get("transaction_date"):
        try:
            for fmt in ["%Y-%m-%d", "%d/%m/%y", "%d-%m-%y", "%d/%m/%Y", "%d-%m-%Y"]:
                try:
                    trans_date = datetime.strptime(expense_dict["transaction_date"], fmt)
                    break
                except:
                    continue
        except:
            pass
    
    expense_dict["transaction_date"] = trans_date
    
    # Generate fingerprint for duplicate detection
    expense_dict["fingerprint"] = generate_fingerprint(
        expense_dict["amount"],
        expense_dict["transaction_type"],
        trans_date
    )
    
    expense_obj = Expense(**expense_dict)
    await db.expenses.insert_one(expense_obj.dict())
    return expense_obj

@api_router.get("/expenses", response_model=List[Expense])
async def get_expenses(
    category_id: Optional[str] = None,
    transaction_type: Optional[str] = None,
    limit: int = 100
):
    query = {}
    if category_id:
        query["category_id"] = category_id
    if transaction_type:
        query["transaction_type"] = transaction_type
    
    expenses = await db.expenses.find(query).sort("created_at", -1).to_list(limit)
    return [Expense(**exp) for exp in expenses]

@api_router.get("/expenses/{expense_id}", response_model=Expense)
async def get_expense(expense_id: str):
    expense = await db.expenses.find_one({"id": expense_id})
    if not expense:
        raise HTTPException(status_code=404, detail="Expense not found")
    return Expense(**expense)

@api_router.delete("/expenses/{expense_id}")
async def delete_expense(expense_id: str):
    result = await db.expenses.delete_one({"id": expense_id})
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Expense not found")
    return {"message": "Expense deleted successfully"}

# Summary endpoint
@api_router.get("/expenses/summary/totals")
async def get_expense_summary():
    expenses = await db.expenses.find().to_list(1000)
    
    total_debit = sum(exp["amount"] for exp in expenses if exp.get("transaction_type") == "debit")
    total_credit = sum(exp["amount"] for exp in expenses if exp.get("transaction_type") == "credit")
    
    # Category-wise breakdown
    category_totals = {}
    for exp in expenses:
        cat_id = exp.get("category_id", "others")
        if cat_id not in category_totals:
            category_totals[cat_id] = {"debit": 0, "credit": 0, "count": 0}
        category_totals[cat_id][exp.get("transaction_type", "debit")] += exp["amount"]
        category_totals[cat_id]["count"] += 1
    
    return {
        "total_debit": total_debit,
        "total_credit": total_credit,
        "balance": total_credit - total_debit,
        "total_transactions": len(expenses),
        "category_totals": category_totals
    }

# Keyword Rules CRUD
@api_router.post("/keyword-rules", response_model=KeywordRule)
async def create_keyword_rule(rule: KeywordRuleCreate):
    existing = await db.keyword_rules.find_one({"keyword": rule.keyword.lower()})
    if existing:
        raise HTTPException(status_code=400, detail="Keyword rule already exists")
    
    rule_dict = rule.dict()
    rule_dict["keyword"] = rule_dict["keyword"].lower()
    rule_obj = KeywordRule(**rule_dict)
    await db.keyword_rules.insert_one(rule_obj.dict())
    return rule_obj

@api_router.get("/keyword-rules", response_model=List[KeywordRule])
async def get_keyword_rules():
    rules = await db.keyword_rules.find().to_list(100)
    return [KeywordRule(**rule) for rule in rules]

@api_router.delete("/keyword-rules/{rule_id}")
async def delete_keyword_rule(rule_id: str):
    result = await db.keyword_rules.delete_one({"id": rule_id})
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Keyword rule not found")
    return {"message": "Keyword rule deleted successfully"}

# SMS Parsing
@api_router.post("/parse-sms", response_model=SMSParseResponse)
async def parse_sms(request: SMSParseRequest):
    result = await parse_sms_message(request.sms_body)
    return SMSParseResponse(**result)

# Check duplicate endpoint
@api_router.post("/check-duplicate")
async def check_duplicate_endpoint(amount: float, transaction_type: str, transaction_date: Optional[str] = None):
    trans_date = None
    if transaction_date:
        try:
            for fmt in ["%Y-%m-%d", "%d/%m/%y", "%d-%m-%y"]:
                try:
                    trans_date = datetime.strptime(transaction_date, fmt)
                    break
                except:
                    continue
        except:
            pass
    
    is_dup = await check_duplicate(amount, transaction_type, trans_date)
    return {"is_duplicate": is_dup}

# PDF Upload and Parse
@api_router.post("/upload-pdf")
async def upload_pdf_statement(file: UploadFile = File(...)):
    if not file.filename.lower().endswith('.pdf'):
        raise HTTPException(status_code=400, detail="Only PDF files are allowed")
    
    content = await file.read()
    transactions = await parse_pdf_statement(content)
    
    if not transactions:
        raise HTTPException(status_code=400, detail="Could not extract any transactions from PDF")
    
    # Check for duplicates and prepare results
    results = []
    for trans in transactions:
        trans_date = None
        if trans.get("transaction_date"):
            try:
                for fmt in ["%d/%m/%y", "%d-%m-%y", "%d/%m/%Y", "%d-%m-%Y"]:
                    try:
                        trans_date = datetime.strptime(trans["transaction_date"], fmt)
                        break
                    except:
                        continue
            except:
                pass
        
        is_dup = await check_duplicate(trans["amount"], trans["transaction_type"], trans_date)
        
        results.append({
            **trans,
            "is_duplicate": is_dup,
            "suggested_category_id": "income" if trans["transaction_type"] == "credit" else "others"
        })
    
    return {
        "success": True,
        "total_found": len(results),
        "transactions": results
    }

# Export to Excel
@api_router.get("/export/excel")
async def export_to_excel():
    # Fetch all expenses
    expenses = await db.expenses.find().sort("created_at", -1).to_list(10000)
    categories = await db.categories.find().to_list(100)
    
    # Create category lookup
    cat_lookup = {cat["id"]: cat["name"] for cat in categories}
    for cat in DEFAULT_CATEGORIES:
        if cat["id"] not in cat_lookup:
            cat_lookup[cat["id"]] = cat["name"]
    
    # Create workbook
    wb = Workbook()
    
    # Styles
    header_font = Font(bold=True, color="FFFFFF")
    header_fill_expense = PatternFill(start_color="FF6B6B", end_color="FF6B6B", fill_type="solid")
    header_fill_income = PatternFill(start_color="2ECC71", end_color="2ECC71", fill_type="solid")
    header_fill_summary = PatternFill(start_color="4ECDC4", end_color="4ECDC4", fill_type="solid")
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # === EXPENSES SHEET ===
    ws_expenses = wb.active
    ws_expenses.title = "Expenses"
    
    expense_headers = ["Date", "Description", "Category", "Merchant", "Amount (₹)", "Source"]
    for col, header in enumerate(expense_headers, 1):
        cell = ws_expenses.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill_expense
        cell.border = thin_border
        cell.alignment = Alignment(horizontal='center')
    
    expense_rows = [e for e in expenses if e.get("transaction_type") == "debit"]
    for row_idx, exp in enumerate(expense_rows, 2):
        date_val = exp.get("transaction_date") or exp.get("created_at")
        if isinstance(date_val, datetime):
            date_str = date_val.strftime("%d/%m/%Y")
        else:
            date_str = str(date_val) if date_val else ""
        
        ws_expenses.cell(row=row_idx, column=1, value=date_str).border = thin_border
        ws_expenses.cell(row=row_idx, column=2, value=exp.get("description", "")).border = thin_border
        ws_expenses.cell(row=row_idx, column=3, value=cat_lookup.get(exp.get("category_id"), "Others")).border = thin_border
        ws_expenses.cell(row=row_idx, column=4, value=exp.get("merchant", "")).border = thin_border
        ws_expenses.cell(row=row_idx, column=5, value=exp.get("amount", 0)).border = thin_border
        ws_expenses.cell(row=row_idx, column=6, value=exp.get("source", "manual")).border = thin_border
    
    # Adjust column widths
    ws_expenses.column_dimensions['A'].width = 12
    ws_expenses.column_dimensions['B'].width = 35
    ws_expenses.column_dimensions['C'].width = 15
    ws_expenses.column_dimensions['D'].width = 20
    ws_expenses.column_dimensions['E'].width = 15
    ws_expenses.column_dimensions['F'].width = 10
    
    # === INCOME SHEET ===
    ws_income = wb.create_sheet("Income")
    
    income_headers = ["Date", "Description", "Category", "Source", "Amount (₹)"]
    for col, header in enumerate(income_headers, 1):
        cell = ws_income.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill_income
        cell.border = thin_border
        cell.alignment = Alignment(horizontal='center')
    
    income_rows = [e for e in expenses if e.get("transaction_type") == "credit"]
    for row_idx, exp in enumerate(income_rows, 2):
        date_val = exp.get("transaction_date") or exp.get("created_at")
        if isinstance(date_val, datetime):
            date_str = date_val.strftime("%d/%m/%Y")
        else:
            date_str = str(date_val) if date_val else ""
        
        ws_income.cell(row=row_idx, column=1, value=date_str).border = thin_border
        ws_income.cell(row=row_idx, column=2, value=exp.get("description", "")).border = thin_border
        ws_income.cell(row=row_idx, column=3, value=cat_lookup.get(exp.get("category_id"), "Income")).border = thin_border
        ws_income.cell(row=row_idx, column=4, value=exp.get("source", "manual")).border = thin_border
        ws_income.cell(row=row_idx, column=5, value=exp.get("amount", 0)).border = thin_border
    
    ws_income.column_dimensions['A'].width = 12
    ws_income.column_dimensions['B'].width = 35
    ws_income.column_dimensions['C'].width = 15
    ws_income.column_dimensions['D'].width = 10
    ws_income.column_dimensions['E'].width = 15
    
    # === SUMMARY SHEET ===
    ws_summary = wb.create_sheet("Summary")
    
    summary_headers = ["Metric", "Value"]
    for col, header in enumerate(summary_headers, 1):
        cell = ws_summary.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill_summary
        cell.border = thin_border
    
    total_expense = sum(e.get("amount", 0) for e in expense_rows)
    total_income = sum(e.get("amount", 0) for e in income_rows)
    
    summary_data = [
        ("Total Expenses", f"₹{total_expense:,.2f}"),
        ("Total Income", f"₹{total_income:,.2f}"),
        ("Net Balance", f"₹{total_income - total_expense:,.2f}"),
        ("Total Transactions", len(expenses)),
        ("Expense Transactions", len(expense_rows)),
        ("Income Transactions", len(income_rows)),
        ("Report Generated", datetime.utcnow().strftime("%d/%m/%Y %H:%M")),
    ]
    
    for row_idx, (metric, value) in enumerate(summary_data, 2):
        ws_summary.cell(row=row_idx, column=1, value=metric).border = thin_border
        ws_summary.cell(row=row_idx, column=2, value=value).border = thin_border
    
    # Category breakdown
    ws_summary.cell(row=10, column=1, value="Category Breakdown").font = Font(bold=True)
    ws_summary.cell(row=11, column=1, value="Category").font = Font(bold=True)
    ws_summary.cell(row=11, column=2, value="Expenses").font = Font(bold=True)
    ws_summary.cell(row=11, column=3, value="Income").font = Font(bold=True)
    
    cat_totals = {}
    for exp in expenses:
        cat_id = exp.get("category_id", "others")
        if cat_id not in cat_totals:
            cat_totals[cat_id] = {"debit": 0, "credit": 0}
        cat_totals[cat_id][exp.get("transaction_type", "debit")] += exp.get("amount", 0)
    
    for row_idx, (cat_id, totals) in enumerate(cat_totals.items(), 12):
        ws_summary.cell(row=row_idx, column=1, value=cat_lookup.get(cat_id, cat_id))
        ws_summary.cell(row=row_idx, column=2, value=f"₹{totals['debit']:,.2f}")
        ws_summary.cell(row=row_idx, column=3, value=f"₹{totals['credit']:,.2f}")
    
    ws_summary.column_dimensions['A'].width = 25
    ws_summary.column_dimensions['B'].width = 20
    ws_summary.column_dimensions['C'].width = 15
    
    # Save to bytes
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    filename = f"expense_report_{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}.xlsx"
    
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )

# Include the router in the main app
app.include_router(api_router)

app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)

# Configure logging
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(name)s - %(levelname)s - %(message)s'
)
logger = logging.getLogger(__name__)

@app.on_event("shutdown")
async def shutdown_db_client():
    client.close()
